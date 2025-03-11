import openpyxl as xl
from docx import Document
# 加载院校项目信息的Excel文件
workbook = xl.load_workbook('dream programs.xlsx')
sheet = workbook.active  # 获取活动工作表
#引入计数器给后续生成的word文档命名
count = 0
# 用户输入需要生成的文档编号
selected_numbers = input("请输入需要生成的文档编号：")
selected_numbers = [int(num.strip()) for num in selected_numbers.split(",")]
# 遍历Excel中的行数据
for row in sheet.iter_rows(min_row=2, values_only=True):  # 从第二行开始读取数据，避免标题行
    count += 1
    if count not in selected_numbers:
        continue  # 跳过不需要生成的行
    university, major, program = row
    # 加载word letter template
    doc = Document('letter template.docx')
    # 在Word文档中查找并替换占位符
    for paragraph in doc.paragraphs:
        if '{{university}}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{{university}}', university)
        if '{{major}}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{{major}}', major)
        if '{{program}}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{{program}}', program)

    doc.save(f"output_{count}.docx")
